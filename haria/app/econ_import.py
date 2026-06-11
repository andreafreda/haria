"""Import estratti conto BancoPosta / Postepay.

Logica pura e testabile: dato il contenuto di un export (xlsx o csv) come lista
di righe (liste di celle), riconosce il formato, normalizza i movimenti a
{data (ISO), importo (firmato), descrizione, categoria} e ne calcola un hash per
il dedup. Niente I/O di rete, niente DB qui.

Due profili (vedi docs/economia-domestica-analisi.md §10):
  - BancoPosta: header con "Addebiti"/"Accrediti" separati -> importo = accrediti - addebiti
  - Postepay:   header con unico "Importo" gia' firmato (negativo = uscita)
"""
import datetime
import hashlib
import re

# --- categorizzazione deterministica per parola chiave -------------------
# (descrizione lower -> categoria). Primo match vince. 'trasferimento' marca le
# ricariche interne BancoPosta->Postepay per non confonderle con spese reali.
_CATEGORIE_KW: list[tuple[tuple[str, ...], str]] = [
    (("ricarica postepay", "ricarica carta"), "trasferimento"),
    (("prelievo", "prelevamento", "prel. ", "atm"), "prelievo contanti"),
    (("f24", "pagopa", "pago pa", "tributi", "imposta", "tasse", "bollo", "canone rai"), "tasse"),
    (("amazon", "zalando", "ebay", "aliexpress", "shein"), "shopping"),
    (("glovo", "just eat", "justeat", "deliveroo", "uber eats"), "ristoranti"),
    (("tim ", "vodafone", "windtre", "wind tre", "iliad", "fastweb"), "abbonamenti"),
    (("netflix", "spotify", "disney", "prime video", "dazn", "youtube"), "abbonamenti"),
    (("eni", "q8", "ip ", "esso", "tamoil", "benzina", "carburante", "distributore"), "carburante"),
    (("esselunga", "conad", "coop", "lidl", "eurospin", "carrefour", "pam ",
      "supermercato", "md ", "penny", "decò", "deco "), "alimentari"),
    (("trenitalia", "italo", "atac", "gtt", "autostrad", "telepass", "pedaggio",
      "parcheggio", "bus ", "metro"), "trasporti"),
    (("farmacia", "parafarmacia", "medic", "ticket sanitario", "dottor"), "salute"),
    (("commission", "competenze", "spese tenuta", "imposta di bollo"), "spese bancarie"),
    (("stipendio", "emolumenti", "cedolino", "salario"), "stipendio"),
    (("enel", "eni gas", "a2a", "hera", "acea", "iren", "bolletta", "luce", "gas"), "bollette"),
]


def categorize(descrizione: str, importo: float) -> str:
    """Categoria dedotta dalla descrizione. Entrate non riconosciute -> 'entrate varie',
    uscite non riconosciute -> 'altro'."""
    t = (descrizione or "").lower()
    for needles, cat in _CATEGORIE_KW:
        if any(n in t for n in needles):
            return cat
    return "entrate varie" if importo > 0 else "altro"


def parse_date(s) -> str | None:
    """Normalizza una data a ISO YYYY-MM-DD. Accetta dd/mm/yyyy, dd-mm-yyyy,
    yyyy-mm-dd, o datetime/date. None se non parsabile."""
    if s is None:
        return None
    # oggetti datetime/date (openpyxl puo' restituirli)
    if hasattr(s, "year") and hasattr(s, "month") and hasattr(s, "day"):
        return f"{s.year:04d}-{s.month:02d}-{s.day:02d}"
    t = str(s).strip()
    if not t:
        return None
    t = t.split(" ")[0]  # scarta eventuale orario
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", t)
    if m:
        y, mo, d = m.groups()
        return _iso_valid(int(y), int(mo), int(d))
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$", t)
    if m:
        d, mo, y = m.groups()
        return _iso_valid(int(y), int(mo), int(d))
    return None


def _iso_valid(y: int, mo: int, d: int) -> str | None:
    """ISO solo se la data esiste davvero (scarta 32/02 ecc.)."""
    try:
        return datetime.date(y, mo, d).isoformat()
    except ValueError:
        return None


def _to_float(v) -> float | None:
    """Converte una cella importo (it: '1.234,56' o '-12,50') in float. None se vuota."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("€", "").replace(" ", "")
    if not t:
        return None
    # formato italiano: punto migliaia, virgola decimali
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip()).lower()


def _find_header(rows: list[list], needles: list[str]) -> int | None:
    """Indice della prima riga che contiene tutte le needle (match su celle lower)."""
    for i, row in enumerate(rows):
        cells = [_norm(c) for c in row]
        joined = " | ".join(cells)
        if all(any(n in c for c in cells) or n in joined for n in needles):
            return i
    return None


def _col_index(header: list, needle: str) -> int | None:
    for i, c in enumerate(header):
        if needle in _norm(c):
            return i
    return None


def detect_format(rows: list[list]) -> str | None:
    """'bancoposta' | 'postepay' | None, in base agli header presenti."""
    if _find_header(rows, ["addebiti", "accrediti", "descrizione"]) is not None:
        return "bancoposta"
    if _find_header(rows, ["importo", "descrizione"]) is not None:
        return "postepay"
    return None


def parse(rows: list[list]) -> dict:
    """Parsa le righe grezze. Ritorna {formato, conto, movimenti:[...], scartate:int}.
    Ogni movimento: {data, importo, descrizione, categoria, hash}."""
    fmt = detect_format(rows)
    if fmt is None:
        return {"formato": None, "conto": None, "movimenti": [], "scartate": 0,
                "errore": "Formato non riconosciuto (header BancoPosta/Postepay assenti)."}

    if fmt == "bancoposta":
        h = _find_header(rows, ["addebiti", "accrediti", "descrizione"])
        header = rows[h]
        c_data = _col_index(header, "data contabile")
        if c_data is None:
            c_data = _col_index(header, "data")
        c_add = _col_index(header, "addebiti")
        c_acc = _col_index(header, "accrediti")
        c_desc = _col_index(header, "descrizione")
        conto = "bancoposta"
    else:
        h = _find_header(rows, ["importo", "descrizione"])
        header = rows[h]
        c_data = _col_index(header, "data contabile")
        if c_data is None:
            c_data = _col_index(header, "data")
        c_imp = _col_index(header, "importo")
        c_desc = _col_index(header, "descrizione")
        conto = "postepay"

    movimenti = []
    scartate = 0
    seen_keys: dict[str, int] = {}
    for row in rows[h + 1:]:
        if not row or all((c is None or str(c).strip() == "") for c in row):
            continue
        data = parse_date(row[c_data]) if c_data is not None and c_data < len(row) else None
        desc = str(row[c_desc]).strip() if c_desc is not None and c_desc < len(row) and row[c_desc] is not None else ""
        if fmt == "bancoposta":
            add = _to_float(row[c_add]) if c_add is not None and c_add < len(row) else None
            acc = _to_float(row[c_acc]) if c_acc is not None and c_acc < len(row) else None
            if add is None and acc is None:
                scartate += 1
                continue
            importo = (acc or 0.0) - abs(add or 0.0)
        else:
            importo = _to_float(row[c_imp]) if c_imp is not None and c_imp < len(row) else None
            if importo is None:
                scartate += 1
                continue
        if data is None:
            scartate += 1
            continue
        importo = round(importo, 2)
        categoria = categorize(desc, importo)
        # dedup: due righe REALI identiche nello stesso file devono avere hash
        # diversi (altrimenti la 2a sparisce in import_transazioni). La 1a
        # occorrenza usa l'hash storico (no suffisso) per retrocompatibilità
        # col DB; le successive aggiungono |n.
        key = f"{conto}|{data}|{importo:.2f}|{_norm(desc)}"
        n = seen_keys.get(key, 0)
        seen_keys[key] = n + 1
        h_row = (row_hash(conto, data, importo, desc) if n == 0
                 else hashlib.sha1(f"{key}|{n}".encode("utf-8")).hexdigest())
        movimenti.append({
            "data": data,
            "importo": importo,
            "descrizione": desc,
            "categoria": categoria,
            "hash": h_row,
        })
    return {"formato": fmt, "conto": conto, "movimenti": movimenti, "scartate": scartate}


def row_hash(conto: str, data: str, importo: float, descrizione: str) -> str:
    """Hash stabile per dedup (stesso movimento reimportato non si duplica)."""
    key = f"{conto}|{data}|{importo:.2f}|{_norm(descrizione)}"
    return hashlib.sha1(key.encode("utf-8")).hexdigest()


# --- lettori file (xlsx via openpyxl, csv via stdlib) --------------------

def rows_from_xlsx(content: bytes) -> list[list]:
    import io
    from openpyxl import load_workbook
    # read_only=False: alcuni export Poste hanno dimensioni dichiarate non da A1
    # (es. A2:E720) che in read_only mode fanno saltare la lettura delle righe.
    wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def rows_from_csv(content: bytes) -> list[list]:
    import csv
    import io
    text = content.decode("utf-8-sig", errors="replace")
    # autodetect delimitatore (Poste csv spesso ';')
    sample = text[:2048]
    delim = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def rows_from_file(content: bytes, filename: str) -> list[list]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        return rows_from_csv(content)
    return rows_from_xlsx(content)
